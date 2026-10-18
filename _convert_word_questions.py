import os, re, hashlib, json
from pathlib import Path
from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

root = Path(r'E:\答题系统\已导入考试系统')
word_files = sorted(root.glob('*.docx'))
excel_files = sorted(root.glob('*.xlsx'))

def clean(s):
    return re.sub(r'\s+', ' ', str(s or '').replace('\u200b','').replace('\ufeff','')).strip()

def answer_norm(s, typ):
    s = clean(s).upper().replace('，', ',').replace('、', ',').replace(' ', '')
    if typ == '判断':
        return {'TRUE':'对','FALSE':'错','T':'对','F':'错','Y':'对','N':'错'}.get(s, s)
    return ''.join(sorted(set(re.findall(r'[A-E]', s))))

def fp(q):
    payload = [q['type'], clean(q['question']).lower(), sorted(q['options'].items()), q['answer']]
    return hashlib.sha256(json.dumps(payload, ensure_ascii=False, separators=(',', ':')).encode()).hexdigest()

def existing_fps():
    result = set()
    for path in excel_files:
        if path.name == 'Word题库转换结果.xlsx':
            continue
        try: wb = load_workbook(path, read_only=True, data_only=True)
        except Exception: continue
        for ws in wb.worksheets:
            typ = {'choice':'单选','mchoice':'多选','judgement':'判断'}.get(ws.title.lower())
            if not typ: continue
            rows = ws.iter_rows(values_only=True)
            try: headers = [clean(x).lower() for x in next(rows)]
            except StopIteration: continue
            for row in rows:
                vals = {headers[i]: row[i] if i < len(row) else '' for i in range(len(headers))}
                question = clean(vals.get('question') or vals.get('题目') or vals.get('问题'))
                ans = answer_norm(vals.get('answer') or vals.get('答案'), typ)
                if not question or not ans: continue
                opts = {k: clean(vals.get(k.lower()) or vals.get(k)) for k in 'ABCDE' if clean(vals.get(k.lower()) or vals.get(k))}
                if typ == '判断' and not opts: opts = {'T':'对','F':'错'}
                result.add(fp({'type':typ,'question':question,'answer':ans,'options':opts}))
    return result

def parse_inline(text, typ):
    text = clean(text)
    ans_match = re.search(r'(?:答案|正确答案)\s*[:：]\s*([A-E,、\s对错TFYN]+)', text, re.I)
    ans = answer_norm(ans_match.group(1), typ) if ans_match else ''
    if ans_match: text = text[:ans_match.start()].strip()
    if typ == '判断':
        return text, {}, ans
    marks = list(re.finditer(r'(?<![A-Za-z])([A-E])[\.、．:：]\s*', text, re.I))
    if not marks: return text, {}, ans
    question = text[:marks[0].start()].strip()
    opts = {}
    for i, m in enumerate(marks):
        end = marks[i+1].start() if i+1 < len(marks) else len(text)
        opts[m.group(1).upper()] = clean(text[m.end():end])
    return question, opts, ans

def parse_doc(path):
    doc = Document(path)
    lines = []
    for p in doc.paragraphs:
        t = clean(p.text)
        if t: lines.extend([clean(x) for x in t.splitlines() if clean(x)])
    result = []
    typ = None
    block = None
    def finish():
        nonlocal block
        if not block or not block.get('question') or not block.get('answer'): block = None; return
        q = {'type': block['type'], 'question': clean(block['question']), 'answer': answer_norm(block['answer'], block['type']), 'options': dict(block['options']), 'chapter': path.stem, 'explanation': block.get('explanation','')}
        if q['type'] == '判断' and not q['options']: q['options'] = {'T':'对','F':'错'}
        if q['type'] != '判断' and any(not q['options'].get(k) for k in q['answer']): block = None; return
        if q['answer'] and q['question']: result.append(q)
        block = None
    for line in lines:
        low = line.lower()
        if re.search(r'单项选择|单选题|单选', line): finish(); typ = '单选'; continue
        if re.search(r'多项选择|多选题|多选', line): finish(); typ = '多选'; continue
        if re.search(r'判断题|判断', line) and not re.search(r'判断题.*答案', line): finish(); typ = '判断'; continue
        if not typ: continue
        if re.match(r'^(?:答案|正确答案)\s*[:：]', line, re.I):
            if block: block['answer'] = re.split(r'[:：]', line, maxsplit=1)[1]
            finish(); continue
        if re.match(r'^[A-E][\.、．:：]', line, re.I):
            if block:
                m = re.match(r'^([A-E])[\.、．:：]\s*(.*)$', line, re.I)
                block['options'][m.group(1).upper()] = clean(m.group(2))
            continue
        if typ == '判断' and re.match(r'^(?:对|错|正确|错误)\s*$', line):
            continue
        qtext, opts, ans = parse_inline(line, typ)
        if ans:
            finish()
            block = {'type':typ, 'question':qtext, 'options':opts, 'answer':ans}
            finish()
        elif opts:
            if block: finish()
            block = {'type':typ, 'question':qtext, 'options':opts, 'answer':''}
        elif block and not block['answer']:
            block['question'] += ' ' + line
        elif re.search(r'[？?。)）]$', line) or re.match(r'^\d+[\.、]', line):
            if block: finish()
            block = {'type':typ, 'question':re.sub(r'^\d+[\.、]\s*','',line), 'options':{}, 'answer':''}
    finish()
    return result

def load_total_questions(path):
    result = {'单选': [], '多选': [], '判断': []}
    if not path.exists():
        return result
    wb = load_workbook(path, read_only=True, data_only=True)
    type_by_sheet = {'choice': '单选', 'mchoice': '多选', 'judgement': '判断'}
    for ws in wb.worksheets:
        typ = type_by_sheet.get(ws.title.lower())
        if not typ:
            continue
        rows = ws.iter_rows(values_only=True)
        headers = [clean(x).lower() for x in next(rows, ())]
        for row in rows:
            vals = {headers[i]: row[i] if i < len(row) else '' for i in range(len(headers))}
            question = clean(vals.get('question'))
            answer = answer_norm(vals.get('answer'), typ)
            if not question or not answer:
                continue
            opts = {k: clean(vals.get(k)) for k in 'ABCDE' if clean(vals.get(k))}
            if typ == '判断' and not opts:
                opts = {'T': '对', 'F': '错'}
            result[typ].append({'type': typ, 'question': question, 'answer': answer,
                                'options': opts, 'chapter': clean(vals.get('group')),
                                'explanation': clean(vals.get('attachment'))})
    return result

def save_questions(path, questions):
    wb = Workbook()
    wb.remove(wb.active)
    headers = ['score','level','group','question','answer','A','B','C','D','E','attachment']
    for typ, sheet in [('单选','choice'),('多选','mchoice'),('判断','judgement')]:
        ws = wb.create_sheet(sheet)
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='1769AA')
            cell.alignment = Alignment(horizontal='center')
        for i, q in enumerate(questions.get(typ, []), 2):
            opts = q['options']
            ws.append([1, '', q['chapter'], q['question'], q['answer'], opts.get('A',''), opts.get('B',''), opts.get('C',''), opts.get('D',''), opts.get('E',''), q.get('explanation','')])
            fill = 'FFFFFF' if i % 2 == 0 else 'F7F9FC'
            for cell in ws[i]:
                cell.fill = PatternFill('solid', fgColor=fill)
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        ws.freeze_panes = 'A2'
        for col, width in {'A':8,'B':10,'C':28,'D':70,'E':12,'F':34,'G':34,'H':34,'I':34,'J':34,'K':50}.items():
            ws.column_dimensions[col].width = width
        ws.auto_filter.ref = ws.dimensions
        for row in ws.iter_rows():
            for cell in row:
                cell.border = Border(bottom=Side(style='thin', color='D9DEE7'))
    wb.save(path)

out = root / 'Word题库转换结果.xlsx'
all_questions = load_total_questions(out)
seen = existing_fps()
converted = 0
skipped = 0
for path in word_files:
    questions = parse_doc(path)
    for q in questions:
        key = fp(q)
        if key in seen:
            skipped += 1
            continue
        seen.add(key)
        all_questions[q['type']].append(q)
    converted += len(questions)

save_questions(out, all_questions)

for path in word_files:
    path.unlink()
print(json.dumps({'word_files':len(word_files),'parsed_questions':converted,
                  'total_questions':sum(map(len,all_questions.values())),
                  'duplicates_or_skipped':skipped,
                  'output':str(out)}, ensure_ascii=False))
