from pathlib import Path
from openpyxl import load_workbook
import re
root = Path(r'E:\答题系统\已导入考试系统')
seen = {}
duplicates = []
total = 0
for path in sorted(root.glob('*.xlsx')):
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        print(f'SKIP {path.name}: {exc}')
        continue
    for sheet in workbook.worksheets:
        if sheet.title.lower() not in {'choice', 'mchoice', 'judgement'}:
            continue
        rows = sheet.iter_rows(values_only=True)
        header = [str(value or '').strip().lower() for value in next(rows, ())]
        indexes = {value: index for index, value in enumerate(header)}
        if 'question' not in indexes:
            continue
        for row in rows:
            question = str(row[indexes['question']] or '').strip()
            if not question:
                continue
            answer = str(row[indexes.get('answer', -1)] or '').strip().upper() if 'answer' in indexes else ''
            key = (sheet.title.lower(), re.sub(r'\s+', ' ', question).lower(), answer)
            total += 1
            if key in seen:
                duplicates.append((path.name, sheet.title, seen[key], question[:80]))
            else:
                seen[key] = f'{path.name}/{sheet.title}'
print(f'FILES={len(list(root.glob("*.xlsx")))}')
print(f'QUESTIONS={total}')
print(f'DUPLICATES={len(duplicates)}')
for item in duplicates[:30]:
    print('DUPLICATE=', item)
