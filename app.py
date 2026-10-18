import os, re, json, random, sqlite3, secrets, csv, io
from datetime import datetime
from flask import Flask, request, jsonify, send_from_directory, session, send_file
from functools import wraps
from werkzeug.security import generate_password_hash, check_password_hash
from openpyxl import load_workbook

BASE = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE, '已导入考试系统')
DB_PATH = os.path.join(BASE, 'question_bank.db')
app = Flask(__name__, static_folder='static')
app.secret_key = os.environ.get('FLASK_SECRET_KEY') or secrets.token_hex(32)
app.config.update(MAX_CONTENT_LENGTH=16 * 1024 * 1024, SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE='Lax', SESSION_COOKIE_SECURE=os.environ.get('COOKIE_SECURE') == '1')

SCHEMA = '''
CREATE TABLE IF NOT EXISTS users (id INTEGER PRIMARY KEY AUTOINCREMENT, username TEXT UNIQUE NOT NULL, password_hash TEXT NOT NULL, role TEXT NOT NULL DEFAULT 'student', created_at TEXT);
CREATE TABLE IF NOT EXISTS settings (key TEXT PRIMARY KEY, value TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS questions (
 id INTEGER PRIMARY KEY AUTOINCREMENT, type TEXT NOT NULL, chapter TEXT NOT NULL,
 score REAL DEFAULT 1, level TEXT, question TEXT NOT NULL, answer TEXT NOT NULL,
 options TEXT DEFAULT '{}', explanation TEXT DEFAULT '', source TEXT DEFAULT '',
 created_at TEXT, updated_at TEXT
);
CREATE TABLE IF NOT EXISTS favorites (user_id INTEGER NOT NULL, question_id INTEGER NOT NULL, created_at TEXT, PRIMARY KEY(user_id, question_id), FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE, FOREIGN KEY(question_id) REFERENCES questions(id) ON DELETE CASCADE);
CREATE TABLE IF NOT EXISTS wrong_questions (user_id INTEGER NOT NULL, question_id INTEGER NOT NULL, wrong_count INTEGER DEFAULT 1, last_answer TEXT, updated_at TEXT, PRIMARY KEY(user_id, question_id), FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE, FOREIGN KEY(question_id) REFERENCES questions(id) ON DELETE CASCADE);
CREATE TABLE IF NOT EXISTS records (id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER NOT NULL, mode TEXT, total INTEGER, correct INTEGER, score REAL, question_ids TEXT, answers TEXT, created_at TEXT, FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE);
CREATE TABLE IF NOT EXISTS progress (
 user_id INTEGER NOT NULL, device_id TEXT NOT NULL, state TEXT NOT NULL, updated_at TEXT NOT NULL,
 PRIMARY KEY(user_id, device_id), FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE
);
CREATE TABLE IF NOT EXISTS import_batches (
 id INTEGER PRIMARY KEY AUTOINCREMENT, source TEXT NOT NULL, file_name TEXT DEFAULT '', imported INTEGER DEFAULT 0,
 skipped INTEGER DEFAULT 0, duplicates INTEGER DEFAULT 0, created_at TEXT NOT NULL, admin_id INTEGER,
 FOREIGN KEY(admin_id) REFERENCES users(id) ON DELETE SET NULL
);
'''

def db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute('PRAGMA foreign_keys=ON')
    return conn

def now(): return datetime.now().isoformat(timespec='seconds')
def clean(v): return '' if v is None else str(v).strip()
def normalize_answer(v, qtype):
    s = clean(v).upper().replace('，', ',').replace('、', ',').replace(' ', '')
    if qtype == '判断':
        return {'T':'对','F':'错','TRUE':'对','FALSE':'错','Y':'对','N':'错'}.get(s, s)
    return ''.join(sorted(set(re.findall('[A-E]', s))))

def validate_question_payload(data):
    if not isinstance(data, dict): return None, '题目格式无效'
    qtype, question = clean(data.get('type')), clean(data.get('question'))
    if qtype not in ('单选', '多选', '判断'): return None, '题型无效'
    if not question: return None, '题干不能为空'
    answer = normalize_answer(data.get('answer'), qtype)
    options = data.get('options') if isinstance(data.get('options'), dict) else {}
    options = {str(k).upper(): clean(v) for k, v in options.items() if str(k).upper() in 'ABCDE' and clean(v)}
    if qtype == '判断':
        options = options or {'T': '对', 'F': '错'}
        if answer not in ('对', '错'): return None, '判断题答案必须为对或错'
    else:
        if not answer or any(key not in options for key in answer): return None, '选择题答案必须对应已有选项'
        if qtype == '单选' and len(answer) != 1: return None, '单选题只能有一个答案'
    normalized = dict(data, type=qtype, question=question, answer=answer, options=options)
    return normalized, None

def question_fingerprint(data):
    q, _ = validate_question_payload(data)
    if not q: return None
    return json.dumps([q['type'], q['question'], sorted(q['options'].items()), q['answer']], ensure_ascii=False, separators=(',', ':'))

def row_question(row, headers, qtype, source):
    vals = {headers[i]: row[i] if i < len(row) else None for i in range(len(headers))}
    question = clean(vals.get('question') or vals.get('题目') or vals.get('问题'))
    if not question: return None
    ans = normalize_answer(vals.get('answer') or vals.get('答案'), qtype)
    if not ans: return None
    options = {x: clean(vals.get(x.lower()) or vals.get(x)) for x in 'ABCDE' if clean(vals.get(x.lower()) or vals.get(x))}
    if qtype == '判断' and not options: options = {'T': '对', 'F': '错'}
    explanation = clean(vals.get('attachment') or vals.get('解析') or vals.get('explanation') or vals.get('analysis'))
    chapter = clean(vals.get('chapter') or vals.get('章节')) or clean(vals.get('章节名称')) or os.path.splitext(os.path.basename(source))[0]
    if chapter.lower() == 'undefined': chapter = os.path.splitext(os.path.basename(source))[0]
    return (qtype, chapter, clean(vals.get('score')) or '1', clean(vals.get('level')), question, ans, json.dumps(options, ensure_ascii=False), explanation, os.path.basename(source), now(), now())

def import_excel(clear=False):
    conn = db(); conn.executescript(SCHEMA); conn.execute('BEGIN')
    if clear: conn.execute('DELETE FROM questions')
    count, skipped, duplicates, files = 0, 0, 0, 0; fingerprints = existing_fingerprints(conn)
    for path in sorted([os.path.join(DATA_DIR, x) for x in os.listdir(DATA_DIR) if x.lower().endswith('.xlsx')]):
        files += 1
        try: wb = load_workbook(path, read_only=True, data_only=True)
        except Exception: skipped += 1; continue
        for ws in wb.worksheets:
            qtype = {'choice':'单选','mchoice':'多选','judgement':'判断'}.get(ws.title.lower())
            if not qtype: continue
            rows = ws.iter_rows(values_only=True)
            try: headers = [clean(x).lower() for x in next(rows)]
            except StopIteration: continue
            for row in rows:
                item = row_question(row, headers, qtype, path)
                if not item: skipped += 1; continue
                payload={'type':item[0],'chapter':item[1],'score':item[2],'level':item[3],'question':item[4],'answer':item[5],'options':json.loads(item[6]),'explanation':item[7]}
                normalized, error = validate_question_payload(payload)
                if error: skipped += 1; continue
                fp=question_fingerprint(normalized)
                if not fp or fp in fingerprints: duplicates += 1; continue
                fingerprints.add(fp)
                conn.execute('INSERT INTO questions(type,chapter,score,level,question,answer,options,explanation,source,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?)', (normalized['type'], normalized.get('chapter') or '未分类', normalized.get('score', 1), normalized.get('level', ''), normalized['question'], normalized['answer'], json.dumps(normalized['options'], ensure_ascii=False), normalized.get('explanation', ''), item[8], item[9], item[10]))
                count += 1
    record_import(conn, '服务器题库目录', '', count, skipped, duplicates); conn.commit(); conn.close(); return {'files':files,'imported':count,'skipped':skipped,'duplicates':duplicates}

def ensure_db():
    conn = db()
    conn.execute('BEGIN IMMEDIATE')
    try:
        for table, required in {
            'favorites': {'user_id'}, 'wrong_questions': {'user_id'},
            'records': {'user_id'}, 'progress': {'user_id'},
        }.items():
            exists = conn.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name=?", (table,)).fetchone()
            if exists:
                columns = {row['name'] for row in conn.execute(f'PRAGMA table_info({table})')}
                if not required.issubset(columns):
                    legacy = 'legacy_' + table
                    if conn.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name=?", (legacy,)).fetchone():
                        legacy += '_v1'
                    conn.execute(f'ALTER TABLE {table} RENAME TO {legacy}')
        for statement in SCHEMA.split(';'):
            if statement.strip(): conn.execute(statement)
        conn.commit()
    except Exception:
        conn.rollback(); conn.close(); raise
    if not conn.execute("SELECT 1 FROM users WHERE username='admin'").fetchone():
        conn.execute('INSERT INTO users(username,password_hash,role,created_at) VALUES(?,?,?,?)', ('admin', generate_password_hash(os.environ.get('ADMIN_PASSWORD', 'admin123')), 'admin', now()))
    conn.execute("UPDATE questions SET chapter=COALESCE(NULLIF(source,''),'未分类') WHERE chapter IS NULL OR TRIM(chapter)='' OR LOWER(chapter)='undefined'")
    conn.commit()
    conn.execute("UPDATE questions SET options=? WHERE type='判断' AND (options IS NULL OR options='{}' OR options='')", (json.dumps({'T':'对','F':'错'}, ensure_ascii=False),))
    conn.commit(); conn.close()
ensure_db()

def qdict(row, include_answer=True):
    d = dict(row); d['options'] = json.loads(d.get('options') or '{}')
    d['is_favorite'] = bool(d.pop('favorite', 0)); d['is_wrong'] = bool(d.pop('wrong', 0))
    if not include_answer: d.pop('answer', None); d.pop('explanation', None)
    return d

def existing_fingerprints(conn):
    return {question_fingerprint({'type': r['type'], 'question': r['question'], 'answer': r['answer'], 'options': json.loads(r['options'] or '{}')}) for r in conn.execute('SELECT type,question,answer,options FROM questions')}

def record_import(conn, source, file_name, imported, skipped, duplicates):
    conn.execute('INSERT INTO import_batches(source,file_name,imported,skipped,duplicates,created_at,admin_id) VALUES(?,?,?,?,?,?,?)', (source, file_name, imported, skipped, duplicates, now(), session.get('user_id')))

def login_required(admin_only=False):
    def decorate(fn):
        @wraps(fn)
        def wrapped(*args, **kwargs):
            if not session.get('user_id') or (admin_only and session.get('role') != 'admin'):
                return jsonify({'error': '需要登录或管理员权限'}), 401
            return fn(*args, **kwargs)
        return wrapped
    return decorate

@app.route('/api/auth/me')
def auth_me(): return jsonify({'user': session.get('username'), 'role': session.get('role')})

@app.after_request
def no_cache_static(response):
    if request.path.startswith('/static/'):
        response.headers['Cache-Control'] = 'no-cache, no-store, must-revalidate'
    return response

@app.route('/api/auth/login', methods=['POST'])
def auth_login():
    data=request.get_json(silent=True) or {}; c=db(); user=c.execute('SELECT * FROM users WHERE username=?',(clean(data.get('username')),)).fetchone(); c.close()
    if not user or not check_password_hash(user['password_hash'], clean(data.get('password'))): return jsonify({'error':'用户名或密码错误'}),401
    session.clear(); session['user_id']=user['id']; session['username']=user['username']; session['role']=user['role']
    return jsonify({'user':user['username'],'role':user['role']})

@app.route('/api/auth/change-password', methods=['POST'])
@login_required()
def change_password():
    data = request.get_json(silent=True) or {}
    current = clean(data.get('current_password'))
    new_password = clean(data.get('new_password'))
    c = db()
    user = c.execute('SELECT password_hash FROM users WHERE id=?', (session['user_id'],)).fetchone()
    if not user or not check_password_hash(user['password_hash'], current):
        c.close(); return jsonify({'error': '当前密码错误'}), 400
    if len(new_password) < 12:
        c.close(); return jsonify({'error': '新密码至少12位'}), 400
    c.execute('UPDATE users SET password_hash=? WHERE id=?', (generate_password_hash(new_password), session['user_id']))
    c.commit(); c.close(); session.clear()
    return jsonify({'ok': True, 'reauthenticate': True})

@app.route('/api/auth/register', methods=['POST'])
def auth_register():
    c=db(); admin_exists=c.execute("SELECT 1 FROM users WHERE role='admin'").fetchone(); c.close()
    if not admin_exists:
        return jsonify({'error':'管理员账户尚未初始化'}), 503
    data=request.get_json(silent=True) or {}; username=clean(data.get('username')); password=clean(data.get('password'))
    if not re.fullmatch(r'[A-Za-z0-9_]{3,32}',username) or len(password)<6: return jsonify({'error':'用户名需为3-32位字母、数字或下划线，密码至少6位'}),400
    c=db()
    try: c.execute('INSERT INTO users(username,password_hash,role,created_at) VALUES(?,?,?,?)',(username,generate_password_hash(password),'student',now())); c.commit()
    except sqlite3.IntegrityError: c.close(); return jsonify({'error':'用户名已存在'}),409
    c.close(); return jsonify({'ok':True})

@app.route('/api/auth/logout', methods=['POST'])
def auth_logout(): session.clear(); return jsonify({'ok':True})

@app.route('/')
def index(): return send_from_directory('static', 'index.html')
@app.route('/admin')
def admin(): return send_from_directory('static', 'admin.html')
@app.route('/result')
def result_page(): return send_from_directory('static', 'result.html')
def progress_key():
    return clean(request.headers.get('X-Device-ID'))

@app.route('/api/progress', methods=['GET', 'PUT', 'DELETE'])
@login_required()
def progress():
    device_id = progress_key()
    user_id = session['user_id']
    if not device_id:
        return jsonify({'error': 'X-Device-ID is required'}), 400
    c = db()
    if request.method == 'GET':
        row = c.execute('SELECT state FROM progress WHERE user_id=? AND device_id=?', (user_id, device_id)).fetchone()
        if not row:
            c.close()
            return jsonify({'state': None})
        try:
            state = json.loads(row['state'])
        except json.JSONDecodeError:
            c.close()
            return jsonify({'state': None})
        if isinstance(state, dict) and isinstance(state.get('qs'), list):
            ids = [item.get('id') for item in state['qs'] if isinstance(item, dict) and isinstance(item.get('id'), int)]
            if ids:
                placeholders = ','.join('?' * len(ids))
                existing = {r['id'] for r in c.execute(f'SELECT id FROM questions WHERE id IN ({placeholders})', ids)}
                state['qs'] = [item for item in state['qs'] if isinstance(item, dict) and item.get('id') in existing]
                state['answers'] = {str(key): value for key, value in (state.get('answers') or {}).items() if str(key).isdigit() and int(key) in existing} if isinstance(state.get('answers'), dict) else {}
            else:
                state['qs'] = []
        c.close()
        return jsonify({'state': state})
    if request.method == 'DELETE':
        c.execute('DELETE FROM progress WHERE user_id=? AND device_id=?', (user_id, device_id))
    else:
        state = request.get_json(silent=True)
        if not isinstance(state, dict):
            c.close()
            return jsonify({'error': 'state must be a JSON object'}), 400
        c.execute('INSERT INTO progress(user_id,device_id,state,updated_at) VALUES(?,?,?,?) ON CONFLICT(user_id,device_id) DO UPDATE SET state=excluded.state,updated_at=excluded.updated_at', (user_id, device_id, json.dumps(state, ensure_ascii=False), now()))
    c.commit(); c.close()
    return jsonify({'ok': True})

@app.route('/api/stats')
@login_required(admin_only=True)
def stats():
    c=db(); total=c.execute('SELECT COUNT(*) n FROM questions').fetchone()['n']; types=c.execute('SELECT type,COUNT(*) n FROM questions GROUP BY type').fetchall(); chapters=c.execute('SELECT chapter,COUNT(*) n FROM questions GROUP BY chapter ORDER BY chapter').fetchall(); c.close(); return jsonify({'total':total,'types':{r['type']:r['n'] for r in types},'chapters':[dict(r) for r in chapters]})
@app.route('/api/chapters')
def chapters():
    c=db(); rows=c.execute("SELECT chapter,COUNT(*) count FROM questions WHERE chapter IS NOT NULL AND TRIM(chapter)<>'' AND LOWER(TRIM(chapter))<>'undefined' GROUP BY chapter ORDER BY chapter").fetchall(); c.close(); return jsonify([dict(r) for r in rows])
@app.route('/api/questions')
@login_required(admin_only=True)
def questions():
    c=db(); where=[]; args=[]; search=clean(request.args.get('search')); typ=clean(request.args.get('type')); types=request.args.getlist('types'); chapter=clean(request.args.get('chapter'))
    if search: where.append('(question LIKE ? OR chapter LIKE ?)'); args += [f'%{search}%',f'%{search}%']
    if types: where.append('type IN (%s)' % ','.join('?' * len(types))); args += types
    elif typ: where.append('type=?'); args.append(typ)
    if chapter: where.append('chapter=?'); args.append(chapter)
    try:
        page = max(1, int(request.args.get('page', 1)))
        page_size = min(100, max(1, int(request.args.get('page_size', 50))))
    except (TypeError, ValueError):
        c.close()
        return jsonify({'error': '分页参数无效'}), 400
    offset = (page - 1) * page_size
    condition=' AND '.join(where) if where else '1=1'
    sql='SELECT * FROM questions WHERE '+condition+' ORDER BY id DESC LIMIT ? OFFSET ?'; rows=c.execute(sql,args+[page_size,offset]).fetchall()
    total=c.execute('SELECT COUNT(*) n FROM questions WHERE '+condition,args).fetchone()['n']; c.close()
    return jsonify({'items':[qdict(r) for r in rows],'total':total,'page':page,'page_size':page_size})
@app.route('/api/questions/clear', methods=['POST'])
@login_required(admin_only=True)
def clear_questions():
    c = db()
    count = c.execute('SELECT COUNT(*) n FROM questions').fetchone()['n']
    c.execute('DELETE FROM questions')
    c.commit()
    c.close()
    return jsonify({'ok': True, 'deleted': count})

@app.route('/api/questions/<int:qid>', methods=['POST','PUT','DELETE'])
@login_required(admin_only=True)
def question_detail(qid):
    c=db()
    if request.method=='DELETE': c.execute('DELETE FROM questions WHERE id=?',(qid,)); c.commit(); c.close(); return jsonify({'ok':True})
    data=request.get_json(silent=True) or {}; normalized, error = validate_question_payload(data)
    if error: c.close(); return jsonify({'error': error}), 400
    if request.method == 'PUT' and c.execute('SELECT 1 FROM questions WHERE id=?', (qid,)).fetchone() is None:
        c.close(); return jsonify({'error': '题目不存在'}), 404
    fields=(normalized['type'], clean(normalized.get('chapter')) or '未分类', normalized.get('score',1), clean(normalized.get('level')), normalized['question'], normalized['answer'], json.dumps(normalized['options'],ensure_ascii=False), clean(normalized.get('explanation')), clean(normalized.get('source')) or '手工录入', now(), now())
    if request.method=='POST': c.execute('INSERT INTO questions(type,chapter,score,level,question,answer,options,explanation,source,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?)',fields)
    else: c.execute('UPDATE questions SET type=?,chapter=?,score=?,level=?,question=?,answer=?,options=?,explanation=?,source=?,updated_at=? WHERE id=?',fields+(qid,))
    c.commit(); c.close(); return jsonify({'ok':True})

@app.route('/api/questions/batch-delete', methods=['POST'])
@login_required(admin_only=True)
def batch_delete():
    ids = (request.get_json(silent=True) or {}).get('ids', [])
    if not isinstance(ids, list) or not ids or len(ids) > 200 or not all(isinstance(x, int) for x in ids): return jsonify({'error':'ids数量无效'}), 400
    c=db(); c.execute('BEGIN'); c.executemany('DELETE FROM questions WHERE id=?', [(x,) for x in set(ids)]); deleted=c.total_changes; c.commit(); c.close(); return jsonify({'ok':True,'deleted':deleted})

@app.route('/api/questions/batch-update', methods=['POST'])
@login_required(admin_only=True)
def batch_update():
    data=request.get_json(silent=True) or {}; ids=data.get('ids', []); chapter=clean(data.get('chapter'))
    if not isinstance(ids,list) or not ids or len(ids)>200 or not chapter: return jsonify({'error':'ids或章节无效'}), 400
    c=db(); c.execute('BEGIN'); c.executemany('UPDATE questions SET chapter=?,updated_at=? WHERE id=?', [(chapter,now(),x) for x in set(ids)]); updated=c.total_changes; c.commit(); c.close(); return jsonify({'ok':True,'updated':updated})
@app.route('/api/import', methods=['POST'])
@login_required(admin_only=True)
def do_import(): return jsonify(import_excel(bool((request.json or {}).get('clear'))))

@app.route('/api/import-mapped', methods=['POST'])
@login_required(admin_only=True)
def import_mapped():
    data = request.get_json(silent=True) or {}; questions = data.get('questions')
    if not isinstance(questions, list): return jsonify({'error': 'questions must be a list'}), 400
    c = db(); imported = skipped = duplicates = 0; fingerprints = existing_fingerprints(c); c.execute('BEGIN')
    for item in questions:
        normalized, error = validate_question_payload(item)
        if error: skipped += 1; continue
        fp=question_fingerprint(normalized)
        if fp in fingerprints: duplicates += 1; continue
        fingerprints.add(fp)
        c.execute('INSERT INTO questions(type,chapter,score,level,question,answer,options,explanation,source,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?)', (normalized['type'], clean(normalized.get('chapter')) or '未分类', normalized.get('score',1), clean(normalized.get('level')), normalized['question'], normalized['answer'], json.dumps(normalized['options'], ensure_ascii=False), clean(normalized.get('explanation')), '上传Excel（字段映射）', now(), now())); imported += 1
    record_import(c, '上传Excel（字段映射）', clean(data.get('file_name')), imported, skipped, duplicates); c.commit(); c.close()
    return jsonify({'imported': imported, 'skipped': skipped, 'duplicates': duplicates})

@app.route('/api/upload-import', methods=['POST'])
@login_required(admin_only=True)
def upload_import():
    files = request.files.getlist('files')
    chapter = clean(request.form.get('chapter'))
    count = skipped = duplicates = 0
    file_names = []
    c = db()
    fingerprints = existing_fingerprints(c)

    try:
        c.execute('BEGIN')
        for file in files:
            filename = clean(file.filename)
            if not filename or not filename.lower().endswith('.xlsx'):
                skipped += 1
                continue
            file_names.append(filename)
            try:
                wb = load_workbook(file, read_only=True, data_only=True)
            except Exception:
                skipped += 1
                continue

            for ws in wb.worksheets:
                qtype = {'choice': '单选', 'mchoice': '多选', 'judgement': '判断'}.get(ws.title.lower())
                if not qtype:
                    continue
                rows = ws.iter_rows(values_only=True)
                try:
                    headers = [clean(x).lower() for x in next(rows)]
                except StopIteration:
                    continue

                for row in rows:
                    item = row_question(row, headers, qtype, filename)
                    if not item:
                        skipped += 1
                        continue
                    has_chapter = (
                        clean(row[headers.index('chapter')])
                        if 'chapter' in headers and headers.index('chapter') < len(row)
                        else clean(row[headers.index('章节')])
                        if '章节' in headers and headers.index('章节') < len(row)
                        else ''
                    )
                    if chapter and not has_chapter:
                        item = (item[0], chapter) + item[2:]
                    payload = {
                        'type': item[0], 'chapter': item[1], 'score': item[2],
                        'level': item[3], 'question': item[4], 'answer': item[5],
                        'options': json.loads(item[6]), 'explanation': item[7]
                    }
                    normalized, error = validate_question_payload(payload)
                    if error:
                        skipped += 1
                        continue
                    fingerprint = question_fingerprint(normalized)
                    if not fingerprint or fingerprint in fingerprints:
                        duplicates += 1
                        continue
                    fingerprints.add(fingerprint)
                    c.execute(
                        'INSERT INTO questions('
                        'type,chapter,score,level,question,answer,options,'
                        'explanation,source,created_at,updated_at'
                        ') VALUES(?,?,?,?,?,?,?,?,?,?,?)',
                        (
                            normalized['type'],
                            normalized.get('chapter') or '未分类',
                            normalized.get('score', 1),
                            normalized.get('level', ''),
                            normalized['question'],
                            normalized['answer'],
                            json.dumps(normalized['options'], ensure_ascii=False),
                            normalized.get('explanation', ''),
                            item[8], item[9], item[10]
                        )
                    )
                    count += 1

        record_import(
            c, '上传Excel', ', '.join(file_names), count, skipped, duplicates
        )
        c.commit()
    except Exception:
        c.rollback()
        c.close()
        raise
    c.close()
    return jsonify({
        'files': len(files),
        'imported': count,
        'skipped': skipped,
        'duplicates': duplicates
    })
@app.route('/api/quiz', methods=['POST'])
@login_required()
def quiz():
    d=request.json or {}; types=d.get('types') or []
    counts=d.get('counts')
    base=[]; args=[]
    if counts is None:
        n=max(1,min(int(d.get('count',10)),100))
        if types: base.append('type IN (%s)'%','.join('?'*len(types))); args += types
        if d.get('chapter'): base.append('chapter=?'); args.append(d['chapter'])
        if d.get('wrong'): base.append('id IN (SELECT question_id FROM wrong_questions WHERE user_id=?)'); args.append(session['user_id'])
        if d.get('favorite'): base.append('id IN (SELECT question_id FROM favorites WHERE user_id=?)'); args.append(session['user_id'])
        c=db(); rows=c.execute('SELECT * FROM questions WHERE '+(' AND '.join(base) if base else '1=1')+' ORDER BY RANDOM() LIMIT ?',args+[n]).fetchall(); c.close(); return jsonify([qdict(r,False) for r in rows])
    requested = {t: max(0, min(int(counts.get(t, 0)), 100)) for t in ('单选','多选','判断')}
    if d.get('chapter'): base.append('chapter=?'); args.append(d['chapter'])
    if d.get('wrong'): base.append('id IN (SELECT question_id FROM wrong_questions WHERE user_id=?)'); args.append(session['user_id'])
    if d.get('favorite'): base.append('id IN (SELECT question_id FROM favorites WHERE user_id=?)'); args.append(session['user_id'])
    c=db(); rows=[]
    for typ, amount in requested.items():
        if amount <= 0: continue
        where=base + ['type=?']; query_args=args + [typ]
        rows.extend(c.execute('SELECT * FROM questions WHERE '+ ' AND '.join(where) +' ORDER BY RANDOM() LIMIT ?', query_args+[amount]).fetchall())
    c.close(); random.shuffle(rows); return jsonify([qdict(r,False) for r in rows])
@app.route('/api/submit', methods=['POST'])
@login_required()
def submit():
    d=request.json or {}; answers=d.get('answers',{}); ids=[int(x) for x in d.get('question_ids',[])]; c=db(); result=[]; correct=0
    for qid in ids:
        r=c.execute('SELECT * FROM questions WHERE id=?',(qid,)).fetchone()
        if not r: continue
        got=normalize_answer(answers.get(str(qid),answers.get(qid,'')),r['type']); ok=got==r['answer']; correct += ok
        if not ok: c.execute('INSERT INTO wrong_questions(user_id,question_id,wrong_count,last_answer,updated_at) VALUES(?,?,?,?,?) ON CONFLICT(user_id,question_id) DO UPDATE SET wrong_count=wrong_count+1,last_answer=excluded.last_answer,updated_at=excluded.updated_at',(session['user_id'],qid,1,got,now()))
        else: c.execute('DELETE FROM wrong_questions WHERE user_id=? AND question_id=?', (session['user_id'], qid))
        result.append({'id':qid,'correct':ok,'answer':r['answer'],'explanation':r['explanation'],'user_answer':got})
    score=round(sum(float(c.execute('SELECT score FROM questions WHERE id=?',(x,)).fetchone()['score'] or 1) for x in ids if any(z['id']==x and z['correct'] for z in result)),2)
    c.execute('INSERT INTO records(user_id,mode,total,correct,score,question_ids,answers,created_at) VALUES(?,?,?,?,?,?,?,?)',(session['user_id'],d.get('mode','练习'),len(ids),correct,score,json.dumps(ids),json.dumps(answers,ensure_ascii=False),now())); c.commit(); c.close(); return jsonify({'total':len(ids),'correct':correct,'score':score,'results':result})
@app.route('/api/favorites/<int:qid>', methods=['POST','DELETE'])
@login_required()
def favorite(qid):
    c=db(); user_id=session['user_id']; c.execute('DELETE FROM favorites WHERE user_id=? AND question_id=?',(user_id,qid));
    if request.method=='POST': c.execute('INSERT INTO favorites(user_id,question_id,created_at) VALUES(?,?,?)',(user_id,qid,now()))
    c.commit(); c.close(); return jsonify({'ok':True,'favorite':request.method=='POST'})
@app.route('/api/review/<kind>')
@login_required()
def review(kind):
    if kind not in ('wrong','favorite'): return jsonify([]),400
    table='wrong_questions' if kind=='wrong' else 'favorites'; order_col='updated_at' if kind=='wrong' else 'created_at'; c=db(); rows=c.execute(f'SELECT q.*,1 as {kind} FROM questions q JOIN {table} x ON x.question_id=q.id WHERE x.user_id=? ORDER BY x.{order_col} DESC',(session['user_id'],)).fetchall(); c.close(); return jsonify([qdict(r,False) for r in rows])
@app.route('/api/questions/export')
@login_required(admin_only=True)
def export_questions():
    fmt=clean(request.args.get('format')).lower() or 'json'; c=db(); where=[]; args=[]
    search=clean(request.args.get('search')); chapter=clean(request.args.get('chapter')); types=request.args.getlist('types')
    if search: where.append('(question LIKE ? OR chapter LIKE ?)'); args += [f'%{search}%', f'%{search}%']
    if chapter: where.append('chapter=?'); args.append(chapter)
    if types: where.append('type IN (%s)' % ','.join('?'*len(types))); args += types
    rows=c.execute('SELECT * FROM questions WHERE '+(' AND '.join(where) if where else '1=1')+' ORDER BY id', args).fetchall(); c.close()
    data=[qdict(r) for r in rows]
    if fmt == 'json': return jsonify(data)
    if fmt != 'csv': return jsonify({'error':'format必须为csv或json'}),400
    output=io.StringIO(); writer=csv.writer(output); writer.writerow(['id','type','chapter','score','level','question','answer','options','explanation','source','created_at','updated_at'])
    for item in data: writer.writerow([item.get(k,'') if k != 'options' else json.dumps(item['options'],ensure_ascii=False) for k in ['id','type','chapter','score','level','question','answer','options','explanation','source','created_at','updated_at']])
    return send_file(io.BytesIO(output.getvalue().encode('utf-8-sig')), mimetype='text/csv', as_attachment=True, download_name='questions.csv')

@app.route('/api/questions/backup', methods=['GET'])
@login_required(admin_only=True)
def backup_questions():
    c=db(); rows=c.execute('SELECT * FROM questions ORDER BY id').fetchall(); c.close()
    return send_file(io.BytesIO(json.dumps([qdict(r) for r in rows], ensure_ascii=False, indent=2).encode('utf-8')), mimetype='application/json', as_attachment=True, download_name='questions-backup.json')

@app.route('/api/questions/restore', methods=['POST'])
@login_required(admin_only=True)
def restore_questions():
    data=request.get_json(silent=True)
    items=data if isinstance(data,list) else (data or {}).get('questions') if isinstance(data,dict) else None
    if not isinstance(items,list): return jsonify({'error':'备份结构无效'}),400
    normalized=[]
    for item in items:
        q,error=validate_question_payload(item)
        if error: return jsonify({'error':'备份包含无效题目：'+error}),400
        normalized.append(q)
    backup_dir=os.path.join(BASE,'backups'); os.makedirs(backup_dir,exist_ok=True)
    c=db(); current=[qdict(r) for r in c.execute('SELECT * FROM questions ORDER BY id').fetchall()]
    backup_path=os.path.join(backup_dir, 'questions-before-restore-'+datetime.now().strftime('%Y%m%d%H%M%S')+'.json')
    with open(backup_path,'w',encoding='utf-8') as f: json.dump(current,f,ensure_ascii=False,indent=2)
    try:
        c.execute('BEGIN'); c.execute('DELETE FROM questions')
        for q in normalized:
            c.execute('INSERT INTO questions(type,chapter,score,level,question,answer,options,explanation,source,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?)',(q['type'],clean(q.get('chapter')) or '未分类',q.get('score',1),clean(q.get('level')),q['question'],q['answer'],json.dumps(q['options'],ensure_ascii=False),clean(q.get('explanation')),clean(q.get('source')) or '备份恢复',now(),now()))
        c.commit()
    except Exception:
        c.rollback(); c.close(); raise
    c.close(); return jsonify({'ok':True,'restored':len(normalized),'backup':backup_path})

@app.route('/api/import-history')
@login_required(admin_only=True)
def import_history():
    c=db(); rows=c.execute('SELECT b.*,u.username AS admin_username FROM import_batches b LEFT JOIN users u ON u.id=b.admin_id ORDER BY b.id DESC LIMIT 100').fetchall(); c.close(); return jsonify([dict(r) for r in rows])

@app.route('/api/admin/overview-stats')
@login_required(admin_only=True)
def overview_stats():
    c=db(); row=c.execute("SELECT COUNT(*) AS students, COALESCE(SUM(total),0) AS total_questions, COALESCE(SUM(correct),0) AS total_correct, COUNT(r.id) AS attempts FROM users u LEFT JOIN records r ON r.user_id=u.id WHERE u.role='student'").fetchone(); c.close(); return jsonify({'students':row['students'],'attempts':row['attempts'],'total_questions':row['total_questions'],'average_accuracy':round(row['total_correct']/row['total_questions']*100,2) if row['total_questions'] else 0})

@app.route('/api/admin/student-stats')
@login_required(admin_only=True)
def student_stats():
    c=db(); rows=c.execute("SELECT u.id,u.username,COUNT(r.id) attempts,COALESCE(SUM(r.total),0) total_questions,COALESCE(SUM(r.correct),0) total_correct FROM users u LEFT JOIN records r ON r.user_id=u.id WHERE u.role='student' GROUP BY u.id ORDER BY u.username").fetchall(); c.close(); return jsonify([dict(r, accuracy=round(r['total_correct']/r['total_questions']*100,2) if r['total_questions'] else 0) for r in rows])

@app.route('/api/admin/users')
@login_required(admin_only=True)
def admin_users():
    c=db(); rows=c.execute("SELECT id,username,role,created_at FROM users ORDER BY role DESC,username").fetchall(); c.close(); return jsonify([dict(r) for r in rows])

@app.route('/api/admin/users/<int:user_id>/reset-password', methods=['POST'])
@login_required(admin_only=True)
def reset_password(user_id):
    password=clean((request.get_json(silent=True) or {}).get('password'))
    if len(password)<6: return jsonify({'error':'密码至少6位'}),400
    c=db(); user=c.execute('SELECT role FROM users WHERE id=?',(user_id,)).fetchone()
    if not user: c.close(); return jsonify({'error':'用户不存在'}),404
    if user['role']=='admin': c.close(); return jsonify({'error':'禁止重置管理员密码'}),400
    c.execute('UPDATE users SET password_hash=? WHERE id=?',(generate_password_hash(password),user_id)); c.commit(); c.close(); return jsonify({'ok':True})

@app.route('/api/records')
@login_required(admin_only=True)
def records():
    c=db(); rows=c.execute('SELECT r.*,u.username FROM records r JOIN users u ON u.id=r.user_id ORDER BY r.id DESC LIMIT 50').fetchall(); c.close(); return jsonify([dict(r) for r in rows])
if __name__=='__main__': app.run(host=os.environ.get('HOST','127.0.0.1'),port=int(os.environ.get('PORT',5000)),debug=False)
